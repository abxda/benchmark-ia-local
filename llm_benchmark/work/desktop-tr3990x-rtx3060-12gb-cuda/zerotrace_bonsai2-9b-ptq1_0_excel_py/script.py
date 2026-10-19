import pandas as pd
import openpyxl
from openpyxl import Workbook

# Leer el CSV
df = pd.read_csv('ventas.csv')

# Calcular ventas totales por region
ventas_por_region = df.groupby('region')['ventas'].sum()

# Crear el libro Excel
wb = Workbook()
ws = wb.active
ws.title = 'Resumen'

# Crear encabezados
ws['A1'] = 'region'
ws['B1'] = 'ventas_totales'

# Agregar filas con datos
for row_idx, (region, total) in enumerate(ventas_por_region.items(), 2):
    ws.cell(row=row_idx, column=1, value=region)
    ws.cell(row=row_idx, column=2, value=total)

# Guardar el archivo
wb.save('reporte.xlsx')

print("Archivo 'reporte.xlsx' creado exitosamente.")